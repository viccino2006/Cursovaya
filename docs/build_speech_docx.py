"""build_speech_docx.py — конвертирует speech.md в speech.docx, удобный для печати."""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt, Mm

HERE = Path(__file__).parent
SRC = HERE / "speech.md"
OUT = HERE / "speech.docx"


def main() -> None:
    text = SRC.read_text(encoding="utf-8")
    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Mm(25)
        section.right_margin = Mm(15)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(stripped[2:].strip())
            r.font.bold = True
            r.font.size = Pt(16)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            r = p.add_run(stripped[3:].strip())
            r.font.bold = True
            r.font.size = Pt(14)
        elif stripped.startswith("---"):
            doc.add_paragraph("")
        elif stripped.startswith("- "):
            # bullet list item; render as plain paragraph with bullet
            content = stripped[2:]
            # convert **bold** patterns
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            r = p.add_run("•  ")
            r.font.size = Pt(13)
            _add_runs_with_bold(p, content)
        elif stripped.startswith("**") and stripped.endswith("**") and stripped.count("**") == 2:
            p = doc.add_paragraph()
            r = p.add_run(stripped[2:-2])
            r.font.bold = True
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.0)
            p.paragraph_format.line_spacing = 1.3
            _add_runs_with_bold(p, stripped)
        i += 1

    doc.save(OUT)
    size_kb = OUT.stat().st_size / 1024
    print(f"Saved {OUT} ({size_kb:.1f} KB)")


def _add_runs_with_bold(p, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.font.bold = True
        else:
            r = p.add_run(part)
        r.font.size = Pt(13)


if __name__ == "__main__":
    main()
